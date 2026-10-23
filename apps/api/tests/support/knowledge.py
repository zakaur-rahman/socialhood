"""Knowledge test support: small PDF and DOCX files built in memory, and a fake DNS for web-page
sources (tests never reach a real website: patch ``ssrf.resolve_host`` with ``fake_dns``)."""

from __future__ import annotations

import io
from collections.abc import Awaitable, Callable
from typing import Any

import docx
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

PUBLIC_IP = "93.184.216.34"
PRIVATE_IP = "10.0.0.7"
HOSTS = {
    "maple.example": [PUBLIC_IP],
    "www.maple.example": [PUBLIC_IP],
    "intranet.maple.example": [PRIVATE_IP],
}


def make_pdf(*pages: list[str]) -> bytes:
    """A PDF with one page per list of lines (Helvetica text pypdf can extract)."""
    writer = PdfWriter()
    for lines in pages:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        ops = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({escaped}) Tj T*")
        ops.append("ET")
        stream = DecodedStreamObject()
        stream.set_data("\n".join(ops).encode("latin-1"))
        page.replace_contents(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx(blocks: list[tuple[str, str]], table: list[list[str]] | None = None) -> bytes:
    """A Word document: ("h1", text) headings and ("p", text) paragraphs, then a table."""
    document = docx.Document()
    for kind, text in blocks:
        if kind == "h1":
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def fake_dns(table: dict[str, list[str]]) -> Callable[[str, int], Awaitable[list[str]]]:
    async def resolve(host: str, port: int) -> list[str]:
        if host not in table:
            raise OSError(f"no such host {host}")
        return table[host]

    return resolve


# ---------------------------------------------------------------- a fixture website (GR-35)

HTML_HEADERS = {"content-type": "text/html; charset=utf-8"}


def article(title: str, topic: str, links: tuple[str, ...] = ()) -> bytes:
    """A page with enough text for the extractor: a heading, two paragraphs about ``topic`` and
    the given links."""
    anchors = "".join(f'<a href="{link}">{link}</a> ' for link in links)
    return (
        f"<html><head><title>{title}</title></head><body><article><h1>{title}</h1>"
        f"<p>{topic}. Our team in Pune answers questions about {topic.lower()} every day of the "
        f"week, and we keep this page up to date.</p>"
        f"<p>If anything about {topic.lower()} is unclear, message us and we will explain it "
        f"in plain words, whenever you ask.</p></article><nav>{anchors}</nav></body></html>"
    ).encode()


def serve(path: str, body: bytes | str, *, status: int = 200, **headers: str) -> Any:
    """Answer GET https://maple.example{path} (the pinned address of maple.example) with
    ``body``; call inside a respx.mock test. Returns the route."""
    import respx

    content = body.encode() if isinstance(body, str) else body
    kind = headers.pop("content_type", "text/html; charset=utf-8")
    return respx.get(f"https://{PUBLIC_IP}{path}").respond(
        status, headers={"content-type": kind, **headers}, content=content
    )


def sitemap_xml(*urls: str) -> str:
    entries = "".join(f"<url><loc>{url}</loc></url>" for url in urls)
    return f'<?xml version="1.0"?><urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">{entries}</urlset>'


def maple_site(*, robots: str | None = None, sitemap: tuple[str, ...] = ()) -> dict[str, Any]:
    """The fixture bakery site: / links to /shipping, /returns, /blog, /private/secret, /menu.pdf
    and another site; /shipping links to /about; the sitemap lists whatever it is given. Returns
    the routes by path; change a page with ``serve`` again (respx uses the newest route)."""
    from typing import cast

    routes: dict[str, Any] = {}
    routes["/robots.txt"] = serve(
        "/robots.txt",
        robots if robots is not None else "User-agent: *\nDisallow: /private\n",
        content_type="text/plain",
    )
    if sitemap:
        routes["/sitemap.xml"] = serve(
            "/sitemap.xml", sitemap_xml(*sitemap), content_type="application/xml"
        )
    else:
        routes["/sitemap.xml"] = serve("/sitemap.xml", "", status=404)
    routes["/"] = serve(
        "/",
        article(
            "Maple Bakery",
            "Fresh cakes and bread",
            (
                "/shipping",
                "/returns",
                "/blog",
                "/private/secret",
                "/menu.pdf",
                "https://other.example/partner",
                "/shipping?utm_source=home",
                "/shipping#top",
            ),
        ),
    )
    routes["/shipping"] = serve(
        "/shipping",
        article("Shipping policy", "Shipping to the UAE and India", ("/about",)),
        etag='"ship1"',
    )
    routes["/returns"] = serve("/returns", article("Returns", "Returns and refunds for cakes"))
    routes["/about"] = serve("/about", article("About us", "Our story and our kitchen"))
    routes["/blog"] = serve("/blog", article("Blog", "News from the bakery", ("/blog/first",)))
    routes["/blog/first"] = serve("/blog/first", article("First post", "Our first blog post"))
    routes["/private/secret"] = serve("/private/secret", article("Secret", "Not for crawlers"))
    return cast("dict[str, Any]", routes)
